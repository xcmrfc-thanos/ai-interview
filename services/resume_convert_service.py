"""Convert uploaded resumes into markdown / txt / pdf preview artifacts.

Upload pipes through this module once; everything downstream (AI analysis,
preview) reads the extracted text and converted artifacts, never the raw doc.
"""

import re
from pathlib import Path

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc", ".txt"}


def convert_resume_files(file_path):
    """Generate same-name .md / .txt next to the source file, and PDF when Word is available.

    Returns dict: {status, text, md_path, txt_path, pdf_path, error}
    """
    path = Path(file_path)
    suffix = path.suffix.lower()
    result = {
        "status": "ok",
        "text": "",
        "md_path": str(path.with_suffix(".md")),
        "txt_path": str(path.with_suffix(".txt")),
        "pdf_path": "",
        "error": "",
    }
    try:
        text = None
        if suffix == ".txt":
            text = path.read_text(encoding="utf-8", errors="replace")
        elif suffix == ".docx":
            text = _docx_to_markdown(path)
            pdf = _word_to_pdf(path)
            result["pdf_path"] = str(pdf) if pdf else ""
        elif suffix == ".doc":
            text = _doc_to_text(path)
            if text:
                pdf = _word_to_pdf(path)
                result["pdf_path"] = str(pdf) if pdf else ""
        elif suffix == ".pdf":
            text = _pdf_to_text(path)
        if text:
            result["text"] = text
            path.with_suffix(".md").write_text(text, encoding="utf-8")
            path.with_suffix(".txt").write_text(_to_plain_text(text), encoding="utf-8")
        else:
            result["status"] = "no-text"
            result["error"] = "未能从中提取到文本内容"
    except Exception as exc:  # noqa: BLE001 - conversion should never break upload
        result["status"] = "failed"
        result["error"] = str(exc)
    return result


def _docx_to_markdown(path):
    from docx import Document

    document = Document(str(path))
    lines = []
    for para in document.paragraphs:
        text = " ".join(para.text.split())
        if not text:
            continue
        style = para.style.name if para.style is not None else ""
        if "Heading 1" in style or "标题 1" in style:
            lines.append(f"# {text}")
        elif "Heading 2" in style or "标题 2" in style:
            lines.append(f"## {text}")
        elif "Heading" in style or "标题" in style:
            lines.append(f"### {text}")
        elif "List" in style and "List Paragraph" not in style or "列表" in style and "列表段落" not in style:
            lines.append(f"- {text}")
        else:
            lines.append(text)
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.replace("|", "/").strip() for cell in row.cells]
            lines.append(f"| {' | '.join(cells)} |")
    return _compact_lines(lines)


def _doc_to_text(path):
    word = _word_com()
    if word is None:
        return None
    try:
        document = word.Documents.Open(str(path), ReadOnly=True)
        text = document.Content.Text
        document.Close(False)
        return _compact_lines(text.splitlines())
    finally:
        word.Quit()


def _word_to_pdf(path):
    word = _word_com()
    if word is None:
        return None
    try:
        document = word.Documents.Open(str(path), ReadOnly=True)
        pdf_path = path.with_suffix(".pdf")
        document.SaveAs2(str(pdf_path), FileFormat=17)
        document.Close(False)
        return pdf_path if pdf_path.exists() else None
    except Exception:
        return None
    finally:
        word.Quit()


def _word_com():
    try:
        import win32com.client
    except ImportError:
        return None
    word = win32com.client.DispatchEx("Word.Application")
    word.Visible = False
    word.DisplayAlerts = 0
    return word


def _pdf_to_text(path):
    from pypdf import PdfReader

    return "\n".join(page.extract_text() or "" for page in PdfReader(str(path)).pages)


def _to_plain_text(markdown):
    lines = []
    for line in markdown.splitlines():
        stripped = line.strip()
        if not stripped:
            lines.append("")
        elif stripped.startswith("|"):
            lines.append(" ".join(cell.strip() for cell in stripped.strip("|").split("|")))
        else:
            lines.append(re.sub(r"^#{1,6}\s+", "", stripped))
    return _compact_lines(lines)


def _compact_lines(lines):
    out = []
    blank = False
    for line in lines:
        text = line.strip()
        if not text:
            if not blank:
                out.append("")
            blank = True
            continue
        blank = False
        out.append(line.strip())
    return "\n".join(out).strip()